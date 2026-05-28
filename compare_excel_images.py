#!/usr/bin/env python3
"""Compare expected image files from Excel against converted images folder.

Excel format expected:
- Column A: product name
- Column C: list with entries containing url=... values, for example:
  [id=1; name=1; url=/media/abc/1.png], [id=2; name=2; url=/media/def/2.png]
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook


URL_PATTERN = re.compile(r"url\s*=\s*([^;\],]+)", re.IGNORECASE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare Excel image URLs (column C) with files in MEDIA-WITHOUT folder."
    )
    parser.add_argument("excel_file", type=Path, help="Path to source Excel file (.xlsx)")
    parser.add_argument(
        "images_folder",
        type=Path,
        nargs="?",
        default=Path("Media-without"),
        help="Folder with converted images (default: Media-without)",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Sheet name (default: active sheet)",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("missing_images_report.csv"),
        help="CSV report path (default: missing_images_report.csv)",
    )
    parser.add_argument(
        "--no-header",
        action="store_true",
        help="Use this if row 1 is data and not headers",
    )
    return parser.parse_args()


def extract_expected_files(column_c_value: str) -> list[str]:
    if not column_c_value:
        return []

    expected = []
    for raw_url in URL_PATTERN.findall(column_c_value):
        clean_url = raw_url.strip().strip("\"'")
        filename = PurePosixPath(clean_url).name
        if filename:
            expected.append(filename)
    return expected


def collect_existing_file_names(images_folder: Path) -> set[str]:
    if not images_folder.exists() or not images_folder.is_dir():
        raise ValueError(f"Images folder not found or not a directory: {images_folder}")

    return {
        image_path.name.lower()
        for image_path in images_folder.rglob("*")
        if image_path.is_file()
    }


def run_comparison(
    excel_file: Path,
    images_folder: Path,
    sheet_name: str | None,
    report_path: Path,
    no_header: bool,
) -> int:
    if not excel_file.exists() or not excel_file.is_file():
        raise ValueError(f"Excel file not found: {excel_file}")

    workbook = load_workbook(excel_file, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    existing_files = collect_existing_file_names(images_folder)
    start_row = 1 if no_header else 2

    rows_with_missing: list[dict[str, str]] = []
    total_rows = 0
    total_expected = 0
    total_missing = 0

    for row_number in range(start_row, worksheet.max_row + 1):
        product_name = worksheet[f"A{row_number}"].value
        raw_media = worksheet[f"C{row_number}"].value

        if product_name is None and raw_media is None:
            continue

        total_rows += 1
        expected_files = extract_expected_files(str(raw_media or ""))
        total_expected += len(expected_files)

        missing_for_row = [
            file_name
            for file_name in expected_files
            if file_name.lower() not in existing_files
        ]

        if missing_for_row:
            total_missing += len(missing_for_row)
            rows_with_missing.append(
                {
                    "row": str(row_number),
                    "product_name": str(product_name or ""),
                    "missing_files": ", ".join(missing_for_row),
                    "expected_count": str(len(expected_files)),
                    "missing_count": str(len(missing_for_row)),
                }
            )

    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", newline="", encoding="utf-8") as csv_file:
        fieldnames = [
            "row",
            "product_name",
            "missing_files",
            "expected_count",
            "missing_count",
        ]
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows_with_missing)

    print(f"Sheet: {worksheet.title}")
    print(f"Rows checked: {total_rows}")
    print(f"Expected images (from column C): {total_expected}")
    print(f"Missing images: {total_missing}")
    print(f"Rows with missing images: {len(rows_with_missing)}")
    print(f"CSV report: {report_path}")

    if rows_with_missing:
        print("\nFirst rows with missing files:")
        for item in rows_with_missing[:10]:
            print(
                f"- Row {item['row']} | Product: {item['product_name']} | "
                f"Missing: {item['missing_files']}"
            )

    return 0


def main() -> int:
    args = parse_args()
    try:
        return run_comparison(
            excel_file=args.excel_file,
            images_folder=args.images_folder,
            sheet_name=args.sheet,
            report_path=args.report,
            no_header=args.no_header,
        )
    except ValueError as exc:
        print(exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
