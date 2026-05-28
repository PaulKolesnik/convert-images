#!/usr/bin/env python3
"""Copy files listed in Excel columns B and C into a new folder, preserving tree.

Expected Excel format:
- Column A: product code/name (optional, not used for matching)
- Column B: one file path (optional)
- Column C: multiple file paths separated by commas (optional)
"""

from __future__ import annotations

import argparse
import csv
import shutil
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Copy images from a large source folder based on Excel column B and C paths, "
            "while preserving original folder structure and file names."
        )
    )
    parser.add_argument("excel_file", type=Path, help="Path to .xlsx file")
    parser.add_argument("source_root", type=Path, help="Root of the large source folder")
    parser.add_argument("output_root", type=Path, help="Destination folder for filtered copy")
    parser.add_argument("--sheet", default=None, help="Sheet name (default: active sheet)")
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("filter_copy_report.csv"),
        help="CSV report path (default: filter_copy_report.csv)",
    )
    parser.add_argument(
        "--no-header",
        action="store_true",
        help="Use this if row 1 is data and not headers",
    )
    return parser.parse_args()


def normalize_excel_path(raw_value: str) -> str:
    text = str(raw_value or "").strip().strip("\"'")
    if not text:
        return ""

    parsed = urlparse(text)
    # Handle values like https://site/media/a/b.png and file:///D:/...
    if parsed.scheme and (parsed.path or parsed.netloc):
        if parsed.scheme.lower() == "file":
            text = parsed.path
        else:
            text = parsed.path

    text = unquote(text)
    return text.strip()


def resolve_source_path(source_root: Path, excel_path_value: str) -> Path:
    normalized = normalize_excel_path(excel_path_value)
    if not normalized:
        return Path()

    candidate = Path(normalized)
    if candidate.is_absolute():
        return candidate

    relative = normalized.lstrip("/\\")
    relative_path = Path(relative.replace("/", "\\"))
    return source_root / relative_path


def extract_paths_from_row(path_b: str, path_c: str) -> list[str]:
    paths: list[str] = []

    b_value = (path_b or "").strip()
    if b_value:
        paths.append(b_value)

    # Column C can contain many paths separated by commas.
    c_raw = (path_c or "").strip()
    if c_raw:
        for part in c_raw.split(","):
            candidate = part.strip()
            if candidate:
                paths.append(candidate)

    return paths


def run_copy(
    excel_file: Path,
    source_root: Path,
    output_root: Path,
    sheet_name: str | None,
    report_path: Path,
    no_header: bool,
) -> int:
    if not excel_file.exists() or not excel_file.is_file():
        raise ValueError(f"Excel file not found: {excel_file}")
    if not source_root.exists() or not source_root.is_dir():
        raise ValueError(f"Source root not found or not a directory: {source_root}")

    source_root = source_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(excel_file, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    start_row = 1 if no_header else 2

    copied_count = 0
    missing_count = 0
    skipped_empty = 0
    empty_paths_count = 0
    seen_sources: set[Path] = set()
    report_rows: list[dict[str, str]] = []

    for row_number in range(start_row, worksheet.max_row + 1):
        product_code_a = str(worksheet[f"A{row_number}"].value or "").strip()
        path_b = str(worksheet[f"B{row_number}"].value or "").strip()
        path_c = str(worksheet[f"C{row_number}"].value or "").strip()

        if not product_code_a and not path_b and not path_c:
            continue

        row_paths = extract_paths_from_row(path_b, path_c)
        if not row_paths:
            skipped_empty += 1
            report_rows.append(
                {
                    "row": str(row_number),
                    "status": "EMPTY_PATH",
                    "source_path": "",
                    "destination_path": "",
                    "note": "Columns B and C are empty or invalid",
                }
            )
            continue

        for row_path in row_paths:
            source_path = resolve_source_path(source_root, row_path)
            if not str(source_path):
                empty_paths_count += 1
                report_rows.append(
                    {
                        "row": str(row_number),
                        "status": "INVALID_PATH",
                        "source_path": "",
                        "destination_path": "",
                        "note": f"Could not parse path: {row_path}",
                    }
                )
                continue

            source_path = source_path.resolve()

            if source_path in seen_sources:
                report_rows.append(
                    {
                        "row": str(row_number),
                        "status": "DUPLICATE",
                        "source_path": str(source_path),
                        "destination_path": "",
                        "note": "Already handled from previous row",
                    }
                )
                continue
            seen_sources.add(source_path)

            if not source_path.exists() or not source_path.is_file():
                missing_count += 1
                report_rows.append(
                    {
                        "row": str(row_number),
                        "status": "MISSING_SOURCE",
                        "source_path": str(source_path),
                        "destination_path": "",
                        "note": "Source file not found",
                    }
                )
                continue

            try:
                relative_path = source_path.relative_to(source_root)
            except ValueError:
                # If path is outside source_root, still copy but isolate under _external.
                relative_path = Path("_external") / source_path.name

            destination_path = output_root / relative_path
            destination_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_path, destination_path)
            copied_count += 1

            report_rows.append(
                {
                    "row": str(row_number),
                    "status": "COPIED",
                    "source_path": str(source_path),
                    "destination_path": str(destination_path),
                    "note": "",
                }
            )

    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("w", newline="", encoding="utf-8") as csv_file:
        fieldnames = ["row", "status", "source_path", "destination_path", "note"]
        writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(report_rows)

    print(f"Sheet: {worksheet.title}")
    print(f"Copied files: {copied_count}")
    print(f"Missing source files: {missing_count}")
    print(f"Rows with empty paths (B and C): {skipped_empty}")
    print(f"Invalid parsed path entries: {empty_paths_count}")
    print(f"Report: {report_path}")

    return 0


def main() -> int:
    args = parse_args()
    try:
        return run_copy(
            excel_file=args.excel_file,
            source_root=args.source_root,
            output_root=args.output_root,
            sheet_name=args.sheet,
            report_path=args.report,
            no_header=args.no_header,
        )
    except ValueError as exc:
        print(exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
